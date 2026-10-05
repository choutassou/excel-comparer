import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pathlib import Path
from tempfile import TemporaryDirectory
import time
import unittest
from unittest.mock import patch

from PySide6.QtCore import QItemSelection, QItemSelectionModel, Qt
from PySide6.QtWidgets import QApplication
from openpyxl import Workbook as ExcelWorkbook
from source.core import CellValue, Change, Sheet, Workbook, compare_workbooks
from source.gui import MainWindow


class GuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.window = MainWindow()
        self.window.show()
        left = Workbook(Path("left.xlsx"), {"main": Sheet("main", {(1, 1): CellValue("before", "s"), (100, 40): CellValue("end", "s")}, 100, 40), "second": Sheet("second")})
        right = Workbook(Path("right.xlsx"), {"main": Sheet("main", {(1, 1): CellValue("after", "s"), (100, 40): CellValue("end", "s")}, 100, 40), "second": Sheet("second")})
        self.window.apply_comparison(compare_workbooks(left, right))
        self.app.processEvents()
        self.addCleanup(self.window.close)

    def test_display_copy_and_read_only(self):
        table = self.window.tables[0]
        index = table.model().index(0, 0)
        self.assertFalse(table.model().flags(index) & Qt.ItemFlag.ItemIsEditable)
        self.assertIsNotNone(index.data(Qt.ItemDataRole.BackgroundRole))
        table.setCurrentIndex(index)
        table.copy_selection()
        self.assertEqual(self.app.clipboard().text(), "before\r\n")
        self.assertEqual(self.window.cell_text.toPlainText(), "before")
        table.selectionModel().select(QItemSelection(index, table.model().index(1, 1)), QItemSelectionModel.SelectionFlag.ClearAndSelect)
        table.copy_selection()
        self.assertEqual(self.app.clipboard().text(), "before\t\r\n\t\r\n")

    def test_sync_navigation_and_tab_position(self):
        left, right = self.window.tables
        left.horizontalScrollBar().setValue(320)
        right.verticalScrollBar().setValue(190)
        self.assertEqual(right.horizontalScrollBar().value(), 320)
        self.assertEqual(left.verticalScrollBar().value(), 190)
        left.setColumnWidth(2, 210)
        self.assertEqual(right.columnWidth(2), 210)
        self.window.tabs.setCurrentIndex(1)
        self.window.tabs.setCurrentIndex(0)
        self.assertEqual(left.horizontalScrollBar().value(), 320)
        self.assertEqual(left.verticalScrollBar().value(), 190)
        self.window.navigate(1)
        self.assertEqual(right.currentIndex().data(), "after")
        self.assertEqual(left.currentIndex().data(), "before")

    def wait_for_load(self):
        deadline = time.monotonic() + 10
        while self.window.worker is not None and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(0.01)
        self.assertIsNone(self.window.worker)

    def test_async_refresh_and_failure(self):
        with TemporaryDirectory() as tmp:
            a, b = Path(tmp) / "a.xlsx", Path(tmp) / "b.xlsx"
            book = ExcelWorkbook()
            book.active["A1"] = "old"
            book.save(a)
            book.active["A1"] = "new"
            book.save(b)
            self.window.load_files(str(a), str(b))
            self.wait_for_load()
            self.assertEqual(self.window.comparison.counts[Change.MODIFIED], 1)
            book.save(a)
            book.close()
            self.window.refresh()
            self.wait_for_load()
            self.assertEqual(self.window.comparison.counts[Change.MODIFIED], 0)
            previous = self.window.comparison
            b.unlink()
            with patch("source.gui.QMessageBox.warning") as warning:
                self.window.refresh()
                self.wait_for_load()
                warning.assert_called_once()
            self.assertIs(self.window.comparison, previous)
            self.assertTrue(self.window.refresh_button.isEnabled())


if __name__ == "__main__":
    unittest.main()
