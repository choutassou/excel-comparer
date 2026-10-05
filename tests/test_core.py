from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook as ExcelWorkbook, load_workbook
from source.core import Change, WorkbookError, compare_files, read_workbook


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.left = Path(self.temp.name) / "left.xlsx"
        self.right = Path(self.temp.name) / "right.xlsx"

    def save(self, path, sheets):
        book = ExcelWorkbook()
        book.remove(book.active)
        for name, cells in sheets.items():
            sheet = book.create_sheet(name)
            for address, value in cells.items():
                sheet[address] = value
        book.save(path)
        book.close()

    def test_two_changes_on_first_and_fifth_sheet(self):
        sheets = {name: {"A1": "same", "B5": "before"} for name in ("main.irai", "main.proc", "main.m_status", "INSERT文", "main.slide")}
        self.save(self.left, sheets)
        for name in ("main.irai", "main.slide"):
            sheets[name]["B5"] = "after"
        self.save(self.right, sheets)
        result = compare_files(self.left, self.right)
        self.assertEqual(result.counts, {Change.DELETED: 0, Change.ADDED: 0, Change.MODIFIED: 2})
        self.assertEqual([len(s.changes) for s in result.sheets], [1, 0, 0, 0, 1])

    def test_blanks_zero_false_and_types(self):
        self.save(self.left, {"main": {"A1": "delete", "C1": 0, "D1": False, "E1": "1", "F1": "", "G1": " "}})
        self.save(self.right, {"main": {"B1": "add", "C1": False, "D1": False, "E1": 1, "G1": " "}})
        self.assertEqual(compare_files(self.left, self.right).counts, {Change.DELETED: 1, Change.ADDED: 1, Change.MODIFIED: 2})

    def test_sheet_union_and_missing_empty_sheet(self):
        self.save(self.left, {"z": {"A1": "x"}, "shared": {"A1": 10}, "empty": {}})
        self.save(self.right, {"shared": {"A1": 10}, "new": {"AA50": 5}})
        result = compare_files(self.left, self.right)
        self.assertEqual([s.name for s in result.sheets], ["z", "shared", "empty", "new"])
        self.assertIsNone(result.sheets[2].right)
        self.assertEqual(result.counts, {Change.DELETED: 1, Change.ADDED: 1, Change.MODIFIED: 0})
        self.assertEqual(result.sheets[-1].columns, 27)
        self.assertEqual(result.sheets[-1].rows, 50)

    def test_formula_refresh_and_no_retained_file_lock(self):
        self.save(self.left, {"main": {"A1": "=1+1"}})
        self.save(self.right, {"main": {"A1": "=2"}})
        initial = compare_files(self.left, self.right)
        self.assertEqual(initial.counts[Change.MODIFIED], 1)
        renamed = self.left.with_name("renamed.xlsx")
        self.left.rename(renamed)
        renamed.unlink()
        self.save(self.left, {"main": {"A1": "=2"}})
        self.assertEqual(compare_files(self.left, self.right).counts[Change.MODIFIED], 0)
        self.assertEqual(initial.left.sheets["main"].cells[(1, 1)].text, "=1+1")

    def test_errors(self):
        with self.assertRaises(WorkbookError):
            read_workbook(self.left)
        with self.assertRaises(WorkbookError):
            read_workbook(self.left.with_suffix(".xls"))
        self.left.write_text("invalid")
        with self.assertRaises(WorkbookError):
            read_workbook(self.left)

    def test_formatting_is_not_a_change(self):
        book = ExcelWorkbook()
        book.active["A1"] = 12
        book.save(self.left)
        book.active["A1"].number_format = "0.00"
        book.save(self.right)
        book.close()
        self.assertFalse(compare_files(self.left, self.right).sheets[0].changes)

    def test_provided_workbooks(self):
        directory = Path(__file__).resolve().parents[1] / "test-data"
        left, right = directory / "data_gen_01.xlsx", directory / "data_gen_02.xlsx"
        if not left.exists() or not right.exists():
            self.skipTest("提供された検証用ファイルがありません")
        result = compare_files(left, right)
        # These user-owned files can change during manual refresh testing.
        # Independently check current contents instead of freezing their counts.
        a, b = load_workbook(left), load_workbook(right)
        try:
            names = list(dict.fromkeys([*a.sheetnames, *b.sheetnames]))
            self.assertEqual([sheet.name for sheet in result.sheets], names)
            for sheet in result.sheets:
                lhs = a[sheet.name] if sheet.name in a.sheetnames else None
                rhs = b[sheet.name] if sheet.name in b.sheetnames else None
                expected = {}
                for row in range(1, max(lhs.max_row if lhs else 0, rhs.max_row if rhs else 0) + 1):
                    for col in range(1, max(lhs.max_column if lhs else 0, rhs.max_column if rhs else 0) + 1):
                        x, y = lhs.cell(row, col) if lhs else None, rhs.cell(row, col) if rhs else None
                        xv, yv = x.value if x else None, y.value if y else None
                        if xv in (None, "") and yv in (None, ""):
                            continue
                        if xv in (None, ""):
                            expected[row, col] = Change.ADDED
                        elif yv in (None, ""):
                            expected[row, col] = Change.DELETED
                        elif (xv, x.data_type) != (yv, y.data_type):
                            expected[row, col] = Change.MODIFIED
                self.assertEqual(sheet.changes, expected)
        finally:
            a.close()
            b.close()


if __name__ == "__main__":
    unittest.main()
