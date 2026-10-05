"""Qt presentation layer. Workbook I/O runs outside the GUI thread."""

import csv
from io import StringIO
from pathlib import Path

from PySide6.QtCore import QAbstractTableModel, QModelIndex, Qt, QThread, Signal
from PySide6.QtGui import QAction, QColor, QKeySequence
from PySide6.QtWidgets import (
    QAbstractItemView, QApplication, QDialog, QDialogButtonBox, QFileDialog,
    QHBoxLayout, QLabel, QLineEdit, QMainWindow, QMenu, QMessageBox, QPushButton,
    QSplitter, QStackedWidget, QTabBar, QTableView, QTextEdit, QVBoxLayout, QWidget,
)
from openpyxl.utils import get_column_letter

from .core import Change, Comparison, SheetComparison, compare_files


COLORS = {
    Change.DELETED: ("#43282f", "#ff929b"),
    Change.ADDED: ("#173d32", "#6fe1ad"),
    Change.MODIFIED: ("#44391f", "#ffd36e"),
}

STYLE = """
QWidget { background: #10151d; color: #dce4ee; font-family: 'Yu Gothic UI', 'Segoe UI'; font-size: 13px; }
QMainWindow, QDialog { background: #10151d; }
QLabel#title { font-size: 24px; font-weight: 700; color: #f2f6fb; }
QLabel#subtitle, QLabel#muted { color: #8899ae; }
QLabel#emptyTitle { font-size: 26px; font-weight: 600; }
QPushButton { background: #202b3b; border: 1px solid #344357; border-radius: 6px; padding: 8px 15px; }
QPushButton:hover { background: #2d3e53; border-color: #758ba7; }
QPushButton:disabled { color: #617087; background: #19212e; }
QPushButton#primary { color: #10151d; background: #efc46b; border-color: #efc46b; font-weight: 700; }
QPushButton#primary:hover { background: #ffda88; }
QLineEdit, QTextEdit { background: #0c1118; border: 1px solid #2c3a4b; border-radius: 4px; padding: 7px; selection-background-color: #405e87; }
QTabBar::tab { padding: 10px 16px; background: #151d28; border-bottom: 2px solid #263243; }
QTabBar::tab:selected { background: #253043; border-bottom: 2px solid #efc46b; }
QTabBar::tab:hover { background: #202b3b; }
QTableView { background: #0e141c; alternate-background-color: #111a25; gridline-color: #202c3b; border: 1px solid #293648; selection-background-color: #365575; selection-color: #ffffff; font-family: 'Consolas', 'Yu Gothic UI'; }
QHeaderView::section { background: #192331; color: #9eb0c6; border: 0; border-right: 1px solid #293648; border-bottom: 1px solid #293648; padding: 5px; }
QTableCornerButton::section { background: #192331; border: 1px solid #293648; }
QScrollBar:horizontal { height: 15px; background: #151d28; }
QScrollBar:vertical { width: 15px; background: #151d28; }
QScrollBar::handle { background: #43546c; border-radius: 5px; min-width: 24px; min-height: 24px; }
QScrollBar::add-line, QScrollBar::sub-line { width: 0; height: 0; }
QScrollBar::add-page, QScrollBar::sub-page { background: #151d28; }
QSplitter::handle { background: #293648; }
QStatusBar { color: #8fa0b6; border-top: 1px solid #293648; }
QMenu { border: 1px solid #344357; }
QMenu::item { padding: 8px 24px; }
QMenu::item:selected { background: #344357; }
"""


class SheetModel(QAbstractTableModel):
    def __init__(self, comparison: SheetComparison, side: str, parent=None):
        super().__init__(parent)
        self.comparison = comparison
        self.sheet = getattr(comparison, side)

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else self.comparison.rows

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else self.comparison.columns

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        address = (index.row() + 1, index.column() + 1)
        value = self.sheet.cells.get(address) if self.sheet else None
        change = self.comparison.changes.get(address)
        if role == Qt.ItemDataRole.DisplayRole:
            return value.text if value else ""
        if role == Qt.ItemDataRole.ToolTipRole:
            return f"{get_column_letter(address[1])}{address[0]}" + (f" · {change.value}" if change else "") + f"\n{value.text if value else '（空白）'}"
        if change and role == Qt.ItemDataRole.BackgroundRole:
            return QColor(COLORS[change][0])
        if change and role == Qt.ItemDataRole.ForegroundRole:
            return QColor(COLORS[change][1])
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole:
            return get_column_letter(section + 1) if orientation == Qt.Orientation.Horizontal else str(section + 1)
        return None


class CellTable(QTableView):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.setSelectionMode(QAbstractItemView.SelectionMode.ContiguousSelection)
        self.setAlternatingRowColors(True)
        self.setWordWrap(False)
        self.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.setVerticalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOn)
        self.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOn)
        self.horizontalHeader().setDefaultSectionSize(126)
        self.horizontalHeader().setMinimumSectionSize(48)
        self.verticalHeader().setDefaultSectionSize(29)
        self.verticalHeader().setFixedWidth(48)
        self.verticalHeader().setSectionResizeMode(self.verticalHeader().ResizeMode.Fixed)
        self.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.customContextMenuRequested.connect(self.context_menu)

    def keyPressEvent(self, event):
        if event.matches(QKeySequence.StandardKey.Copy):
            self.copy_selection()
            event.accept()
        else:
            super().keyPressEvent(event)

    def context_menu(self, position):
        menu = QMenu(self)
        action = menu.addAction("コピー\tCtrl+C", self.copy_selection)
        action.setEnabled(self.selectionModel().hasSelection())
        menu.exec(self.viewport().mapToGlobal(position))

    def copy_selection(self):
        ranges = self.selectionModel().selection()
        if not ranges:
            return
        area = ranges[0]
        if area.width() * area.height() > 1_000_000:
            QMessageBox.information(self, "コピー範囲", "100万セル以下の範囲を選択してください。")
            return
        output = StringIO(newline="")
        writer = csv.writer(output, delimiter="\t", lineterminator="\r\n")
        for row in range(area.top(), area.bottom() + 1):
            writer.writerow([self.model().index(row, col).data() or "" for col in range(area.left(), area.right() + 1)])
        QApplication.clipboard().setText(output.getvalue())


class OpenDialog(QDialog):
    def __init__(self, paths=None, parent=None):
        super().__init__(parent)
        self.setWindowTitle("比較するファイルを開く")
        self.resize(750, 300)
        layout = QVBoxLayout(self)
        title = QLabel("2つのExcelファイルを比較")
        title.setObjectName("title")
        layout.addWidget(title)
        layout.addWidget(QLabel("左を比較元、右を比較先として、同名シート・同じセル位置の値を比較します。"))
        self.fields = []
        for i, label in enumerate(("左側（比較元）", "右側（比較先）")):
            layout.addWidget(QLabel(label))
            row = QHBoxLayout()
            field = QLineEdit(paths[i] if paths else "")
            field.setPlaceholderText("Excelファイルのパス（.xlsx / .xlsm）")
            browse = QPushButton("参照…")
            browse.clicked.connect(lambda checked=False, target=field: self.browse(target))
            row.addWidget(field, 1)
            row.addWidget(browse)
            layout.addLayout(row)
            self.fields.append(field)
        swap = QPushButton("左右を入れ替え")
        swap.clicked.connect(self.swap)
        layout.addWidget(swap, alignment=Qt.AlignmentFlag.AlignLeft)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText("比較")
        buttons.button(QDialogButtonBox.StandardButton.Ok).setObjectName("primary")
        buttons.button(QDialogButtonBox.StandardButton.Cancel).setText("キャンセル")
        buttons.accepted.connect(self.validate)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def browse(self, field):
        path, _ = QFileDialog.getOpenFileName(self, "Excelファイルを選択", field.text(), "Excel (*.xlsx *.xlsm)")
        if path:
            field.setText(path)

    def swap(self):
        left, right = self.paths()
        self.fields[0].setText(right)
        self.fields[1].setText(left)

    def paths(self):
        return tuple(field.text().strip().strip('"') for field in self.fields)

    def validate(self):
        for path in self.paths():
            if not path or not Path(path).is_file():
                QMessageBox.warning(self, "ファイルを確認", "左右それぞれに存在するファイルを指定してください。")
                return
            if Path(path).suffix.lower() not in {".xlsx", ".xlsm"}:
                QMessageBox.warning(self, "対応形式", ".xlsx / .xlsm ファイルを指定してください。")
                return
        self.accept()


class LoadThread(QThread):
    loaded = Signal(object)
    failed = Signal(str)
    progress = Signal(str)

    def __init__(self, paths, parent=None):
        super().__init__(parent)
        self.paths = paths

    def run(self):
        try:
            self.loaded.emit(compare_files(*self.paths, progress=self.progress.emit))
        except Exception as exc:
            self.failed.emit(str(exc))


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Excel比較")
        self.resize(1480, 860)
        self.setMinimumSize(900, 560)
        self.setStyleSheet(STYLE)
        self.comparison: Comparison | None = None
        self.worker = None
        self.current_sheet = None
        self.positions = {}
        self.changes = []
        self.change_index = -1
        self._syncing = False
        root = QWidget()
        self.setCentralWidget(root)
        layout = QVBoxLayout(root)
        layout.setContentsMargins(22, 18, 22, 8)
        layout.setSpacing(12)
        header = QHBoxLayout()
        heading = QVBoxLayout()
        title = QLabel("Excel比較")
        title.setObjectName("title")
        heading.addWidget(title)
        subtitle = QLabel("ワークブックの変更点を、ひと目で。")
        subtitle.setObjectName("subtitle")
        heading.addWidget(subtitle)
        header.addLayout(heading)
        header.addStretch()
        self.open_button = QPushButton("ファイルを開く…")
        self.open_button.setObjectName("primary")
        self.open_button.clicked.connect(self.open_files)
        self.refresh_button = QPushButton("再読み込み  (F5)")
        self.refresh_button.setToolTip("外部で保存された最新の内容を読み込みます（F5）")
        self.refresh_button.clicked.connect(self.refresh)
        self.refresh_button.setEnabled(False)
        header.addWidget(self.open_button)
        header.addWidget(self.refresh_button)
        layout.addLayout(header)
        summary_row = QHBoxLayout()
        self.summary = QLabel("ファイルを選択して比較を開始")
        summary_row.addWidget(self.summary)
        summary_row.addStretch()
        self.previous = QPushButton("‹  前の差分")
        self.next = QPushButton("次の差分  ›")
        self.previous.clicked.connect(lambda: self.navigate(-1))
        self.next.clicked.connect(lambda: self.navigate(1))
        self.previous.setEnabled(False)
        self.next.setEnabled(False)
        summary_row.addWidget(self.previous)
        summary_row.addWidget(self.next)
        layout.addLayout(summary_row)
        self.tabs = QTabBar()
        self.tabs.setExpanding(False)
        self.tabs.setUsesScrollButtons(True)
        self.tabs.setElideMode(Qt.TextElideMode.ElideNone)
        self.tabs.currentChanged.connect(self.show_sheet)
        layout.addWidget(self.tabs)
        self.stack = QStackedWidget()
        empty = QWidget()
        empty_layout = QVBoxLayout(empty)
        empty_layout.addStretch()
        label = QLabel("Excelファイルを並べて比較")
        label.setObjectName("emptyTitle")
        empty_layout.addWidget(label, alignment=Qt.AlignmentFlag.AlignCenter)
        hint = QLabel("削除は赤、追加は緑、変更は黄色で表示します。\n\n.xlsx / .xlsm  ·  表示専用  ·  セルのコピーに対応")
        hint.setObjectName("muted")
        hint.setAlignment(Qt.AlignmentFlag.AlignCenter)
        empty_layout.addWidget(hint)
        empty_layout.addStretch()
        self.stack.addWidget(empty)
        self.splitter = QSplitter(Qt.Orientation.Horizontal)
        self.splitter.setChildrenCollapsible(False)
        self.tables, self.path_labels, self.sheet_labels = [], [], []
        for side in ("左側 · 比較元", "右側 · 比較先"):
            pane = QWidget()
            pane_layout = QVBoxLayout(pane)
            pane_layout.setContentsMargins(0, 0, 0, 0)
            pane_layout.addWidget(QLabel(side))
            path_label = QLineEdit()
            path_label.setReadOnly(True)
            pane_layout.addWidget(path_label)
            sheet_label = QLabel()
            sheet_label.setObjectName("muted")
            pane_layout.addWidget(sheet_label)
            table = CellTable()
            pane_layout.addWidget(table)
            self.path_labels.append(path_label)
            self.sheet_labels.append(sheet_label)
            self.tables.append(table)
            self.splitter.addWidget(pane)
        self.stack.addWidget(self.splitter)
        layout.addWidget(self.stack, 1)
        self.cell_label = QLabel("セルの内容（選択してコピーできます）")
        self.cell_text = QTextEdit()
        self.cell_text.setReadOnly(True)
        self.cell_text.setMaximumHeight(75)
        layout.addWidget(self.cell_label)
        layout.addWidget(self.cell_text)
        for source, target in ((self.tables[0], self.tables[1]), (self.tables[1], self.tables[0])):
            for orientation in ("horizontal", "vertical"):
                bar = getattr(source, orientation + "ScrollBar")()
                other = getattr(target, orientation + "ScrollBar")()
                bar.valueChanged.connect(lambda value, peer=other: self.sync_scroll(peer, value))
            source.horizontalHeader().sectionResized.connect(lambda col, old, size, peer=target: self.sync_width(peer, col, size))
        for key, callback in (("Ctrl+O", self.open_files), ("F5", self.refresh), ("F7", lambda: self.navigate(1)), ("Shift+F7", lambda: self.navigate(-1))):
            action = QAction(self)
            action.setShortcut(QKeySequence(key))
            action.triggered.connect(callback)
            self.addAction(action)
        self.statusBar().showMessage("表示専用  •  同名シートの同じセル位置を比較  •  Ctrl+O: 開く")

    def sync_scroll(self, peer, value):
        if self._syncing:
            return
        self._syncing = True
        peer.setValue(value)
        self._syncing = False

    def sync_width(self, peer, column, size):
        if peer.columnWidth(column) != size:
            peer.setColumnWidth(column, size)

    def open_files(self):
        if self.worker and self.worker.isRunning():
            return
        paths = (str(self.comparison.left.path), str(self.comparison.right.path)) if self.comparison else None
        dialog = OpenDialog(paths, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.load_files(*dialog.paths())

    def refresh(self):
        if self.comparison:
            self.load_files(str(self.comparison.left.path), str(self.comparison.right.path))

    def load_files(self, left, right):
        if self.worker and self.worker.isRunning():
            return
        self.open_button.setEnabled(False)
        self.refresh_button.setEnabled(False)
        self.statusBar().showMessage("Excelファイルを読み込んでいます…")
        self.worker = LoadThread((left, right), self)
        self.worker.progress.connect(self.statusBar().showMessage)
        self.worker.loaded.connect(self.apply_comparison)
        self.worker.failed.connect(self.load_failed)
        self.worker.finished.connect(self.load_finished)
        self.worker.start()

    def load_finished(self):
        self.open_button.setEnabled(True)
        self.refresh_button.setEnabled(self.comparison is not None)
        self.worker.deleteLater()
        self.worker = None

    def load_failed(self, message):
        self.statusBar().showMessage("読み込みに失敗しました。前回の比較結果を表示しています。" if self.comparison else "読み込みに失敗しました。")
        QMessageBox.warning(self, "読み込みエラー", message)

    def save_position(self):
        if self.current_sheet is not None:
            table = self.tables[0]
            self.positions[self.current_sheet] = (table.horizontalScrollBar().value(), table.verticalScrollBar().value())

    def apply_comparison(self, comparison):
        self.save_position()
        previous_name = self.current_sheet
        if self.comparison is None or (self.comparison.left.path, self.comparison.right.path) != (comparison.left.path, comparison.right.path):
            self.positions.clear()
            previous_name = None
        self.current_sheet = None
        self.comparison = comparison
        self.changes = [(i, address) for i, sheet in enumerate(comparison.sheets) for address in sheet.changes]
        self.change_index = -1
        self.previous.setEnabled(bool(self.changes))
        self.next.setEnabled(bool(self.changes))
        self.summary.setText("　　".join(f'<span style="color:{COLORS[kind][1]}">● {count:,} {kind.value}</span>' for kind, count in comparison.counts.items()))
        for label, book in zip(self.path_labels, (comparison.left, comparison.right)):
            label.setText(str(book.path))
            label.setToolTip(str(book.path))
            label.setCursorPosition(0)
        self.tabs.blockSignals(True)
        while self.tabs.count():
            self.tabs.removeTab(0)
        selected = 0
        for i, sheet in enumerate(comparison.sheets):
            suffix = " · 左のみ" if sheet.right is None else " · 右のみ" if sheet.left is None else ""
            self.tabs.addTab(f"{sheet.name}{suffix}" + (f"  ({len(sheet.changes)})" if sheet.changes else ""))
            self.tabs.setTabTextColor(i, QColor("#ffd36e" if sheet.changes or suffix else "#96a8bf"))
            self.tabs.setTabToolTip(i, " / ".join(f"{n} {k.value}" for k, n in sheet.counts.items()) + suffix)
            if sheet.name == previous_name:
                selected = i
        self.tabs.setCurrentIndex(selected)
        self.tabs.blockSignals(False)
        self.stack.setCurrentIndex(1)
        self.show_sheet(selected)

    def show_sheet(self, index):
        if not self.comparison or index < 0 or index >= len(self.comparison.sheets):
            return
        self.save_position()
        sheet = self.comparison.sheets[index]
        self.current_sheet = sheet.name
        self.cell_text.clear()
        self.cell_label.setText("セルの内容（選択してコピーできます）")
        for i, side in enumerate(("left", "right")):
            table = self.tables[i]
            old = table.model()
            old_selection = table.selectionModel()
            table.setModel(SheetModel(sheet, side, table))
            if old:
                old.deleteLater()
            if old_selection:
                old_selection.deleteLater()
            table.selectionModel().currentChanged.connect(lambda current, previous, side_index=i: self.show_cell(side_index, current))
            self.sheet_labels[i].setText("この側にはシートがありません" if getattr(sheet, side) is None else f"{sheet.name}  ·  {len(sheet.changes):,} 箇所の差分")
        for table in self.tables:
            table.doItemsLayout()
            x, y = self.positions.get(sheet.name, (0, 0))
            table.horizontalScrollBar().setValue(x)
            table.verticalScrollBar().setValue(y)
        self.statusBar().showMessage(f"{len(self.comparison.sheets)} シート  •  {len(self.changes):,} 箇所の差分  •  表示専用  •  Ctrl+C: コピー / F5: 再読み込み / F7: 次の差分")

    def show_cell(self, side, index):
        if not index.isValid():
            return
        self.cell_label.setText(f"{'左' if side == 0 else '右'} · {get_column_letter(index.column() + 1)}{index.row() + 1}  —  セルの内容")
        self.cell_text.setPlainText(index.data() or "")

    def navigate(self, direction):
        if not self.changes:
            return
        self.change_index = ((0 if direction > 0 else len(self.changes) - 1) if self.change_index == -1 else (self.change_index + direction) % len(self.changes))
        sheet_index, (row, col) = self.changes[self.change_index]
        self.tabs.setCurrentIndex(sheet_index)
        for table in self.tables:
            index = table.model().index(row - 1, col - 1)
            table.setCurrentIndex(index)
            table.scrollTo(index, QAbstractItemView.ScrollHint.PositionAtCenter)
        self.statusBar().showMessage(f"差分 {self.change_index + 1} / {len(self.changes)}  •  {self.current_sheet} ! {get_column_letter(col)}{row}")

    def closeEvent(self, event):
        if self.worker and self.worker.isRunning():
            self.statusBar().showMessage("読み込み中です。完了後にウィンドウを閉じてください。")
            event.ignore()
        else:
            event.accept()
