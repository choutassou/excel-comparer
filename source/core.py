"""Workbook snapshots and coordinate-based comparison, with no GUI dependencies."""

from dataclasses import dataclass, field
from datetime import date, datetime, time
from enum import Enum
from io import BytesIO
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook


class WorkbookError(Exception):
    """An actionable workbook loading error."""


class Change(str, Enum):
    DELETED = "削除"
    ADDED = "追加"
    MODIFIED = "変更"


@dataclass(frozen=True)
class CellValue:
    value: object
    kind: str

    @property
    def text(self) -> str:
        if isinstance(self.value, bool):
            return "TRUE" if self.value else "FALSE"
        if isinstance(self.value, (date, datetime, time)):
            return self.value.isoformat(sep=" ") if isinstance(self.value, datetime) else self.value.isoformat()
        return str(self.value)


@dataclass
class Sheet:
    name: str
    cells: dict[tuple[int, int], CellValue] = field(default_factory=dict)
    rows: int = 0
    columns: int = 0


@dataclass
class Workbook:
    path: Path
    sheets: dict[str, Sheet]


@dataclass
class SheetComparison:
    name: str
    left: Sheet | None
    right: Sheet | None
    changes: dict[tuple[int, int], Change]

    @property
    def rows(self) -> int:
        return max(self.left.rows if self.left else 0, self.right.rows if self.right else 0, 30)

    @property
    def columns(self) -> int:
        return max(self.left.columns if self.left else 0, self.right.columns if self.right else 0, 12)

    @property
    def counts(self) -> dict[Change, int]:
        return {kind: sum(value == kind for value in self.changes.values()) for kind in Change}


@dataclass
class Comparison:
    left: Workbook
    right: Workbook
    sheets: list[SheetComparison]

    @property
    def counts(self) -> dict[Change, int]:
        totals = dict.fromkeys(Change, 0)
        for sheet in self.sheets:
            for kind, count in sheet.counts.items():
                totals[kind] += count
        return totals


def read_workbook(path: str | Path, progress: Callable[[str], None] | None = None) -> Workbook:
    """Read a detached snapshot. No source file handle survives read_bytes().

    Empty strings and None are blank. Formulas are compared as expressions;
    formatting and cached calculation results do not affect comparison.
    """
    source = Path(path).expanduser().resolve()
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise WorkbookError("対応形式は .xlsx / .xlsm です。.xls は .xlsx に変換してください。")
    book = None
    try:
        with BytesIO(source.read_bytes()) as snapshot:
            book = load_workbook(snapshot, read_only=True, data_only=False, keep_links=False)
            sheets = {}
            for ws in book.worksheets:
                if progress:
                    progress(f"読み込み中: {source.name} / {ws.title}")
                # Avoid expanding massive formatted-but-empty rectangles silently.
                if (ws.max_row or 1) * (ws.max_column or 1) > 5_000_000:
                    raise WorkbookError(f"「{ws.title}」の使用範囲が大きすぎます（上限500万セル）。不要な行・列の書式を削除してください。")
                sheet = Sheet(ws.title)
                scanned = 0
                for row in ws.iter_rows():
                    scanned += len(row)
                    if scanned > 5_000_000:
                        raise WorkbookError(f"「{ws.title}」が読み込み上限500万セルを超えました。")
                    for cell in row:
                        if cell.value is None or cell.value == "":
                            continue
                        value = cell.value
                        if cell.data_type == "f" and not isinstance(value, str):
                            # openpyxl represents array/data-table formulas as objects.
                            value = repr(sorted(vars(value).items()))
                        sheet.cells[(cell.row, cell.column)] = CellValue(value, cell.data_type)
                        sheet.rows = max(sheet.rows, cell.row)
                        sheet.columns = max(sheet.columns, cell.column)
                sheets[ws.title] = sheet
            book.close()
            book = None
            return Workbook(source, sheets)
    except WorkbookError:
        raise
    except PermissionError as exc:
        raise WorkbookError(f"ファイルを読み込めません: {source.name}\nアクセス権限と、他のアプリの排他設定を確認してください。") from exc
    except FileNotFoundError as exc:
        raise WorkbookError(f"ファイルが見つかりません: {source}") from exc
    except Exception as exc:
        raise WorkbookError(f"{source.name} を読み込めません。Excel形式、破損、パスワード保護を確認してください。\n{exc}") from exc
    finally:
        if book is not None:
            book.close()


def compare_workbooks(left: Workbook, right: Workbook) -> Comparison:
    """Match sheet names, then exact cell addresses. Count each cell once."""
    results = []
    for name in dict.fromkeys([*left.sheets, *right.sheets]):
        lhs, rhs = left.sheets.get(name), right.sheets.get(name)
        a = lhs.cells if lhs else {}
        b = rhs.cells if rhs else {}
        changes = {}
        for address in sorted(a.keys() | b.keys()):
            if address not in a:
                changes[address] = Change.ADDED
            elif address not in b:
                changes[address] = Change.DELETED
            elif a[address] != b[address]:
                changes[address] = Change.MODIFIED
        results.append(SheetComparison(name, lhs, rhs, changes))
    return Comparison(left, right, results)


def compare_files(left: str | Path, right: str | Path, progress=None) -> Comparison:
    return compare_workbooks(read_workbook(left, progress), read_workbook(right, progress))
